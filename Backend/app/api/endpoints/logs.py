import io
import logging
from datetime import datetime, timedelta

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy.orm import Session
from sqlalchemy import func
from typing import List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.database.session import get_db, engine
from app.repositories.repositories import log_repo, log_file_repo
from app.schemas.schemas import LogResponse, LogReprocessRequest, LogSummaryResponse
from app.services.watcher import log_watcher_service, parse_source_identity, get_active_root_paths, parser as log_parser
from app.services.rca import auto_trigger
from app.services.gitlab.gitlab_agent import is_gitlab_configured
from app.models.models import Log, LogAnalysis, IncidentDecision, GitlabIssue
import os

logger = logging.getLogger(__name__)

router = APIRouter()

EXPORT_ROW_LIMIT = 10000

# Characters that Excel/Sheets interpret as the start of a formula. Prefixing
# with a single quote prevents formula injection when a log message happens to
# start with one of these (log content is not trusted input) -- openpyxl would
# otherwise store such a string as an actual formula.
_FORMULA_PREFIXES = ("=", "+", "-", "@")

_SOLVED_FILL = PatternFill(start_color="FF1E4620", end_color="FF1E4620", fill_type="solid")   # dark green: error, RCA found
_UNSOLVED_FILL = PatternFill(start_color="FF4A1A1F", end_color="FF4A1A1F", fill_type="solid")  # dark red: error, no RCA yet
_HEADER_FILL = PatternFill(start_color="FF262E42", end_color="FF262E42", fill_type="solid")


def _excel_safe(value: Optional[str]) -> str:
    text = value or ""
    if text and text[0] in _FORMULA_PREFIXES:
        return "'" + text
    return text


def _bucket_expr(granularity: str):
    if engine.dialect.name == "sqlite":
        fmt = "%Y-%m-%d %H:00:00" if granularity == "hour" else "%Y-%m-%d 00:00:00"
        return func.strftime(fmt, Log.timestamp)
    return func.date_trunc(granularity, Log.timestamp)


def _parse_bucket(value, granularity: str) -> datetime:
    if isinstance(value, datetime):
        return value
    fmt = "%Y-%m-%d %H:%M:%S"
    return datetime.strptime(value, fmt)


def _attach_gitlab_issue_info(db: Session, logs: List[Log]) -> List[LogResponse]:
    """
    Bulk-attaches GitLab issue status to each log: whether its incident
    group (any occurrence, not just this exact log) already has an issue
    filed, so the UI can offer "create" vs. "view existing" instead of
    risking a duplicate issue for a recurring problem. One extra query for
    the whole page, not one per row.
    """
    gitlab_configured = is_gitlab_configured(db)
    group_ids = {
        log.analyses[0].incident_group_id
        for log in logs
        if log.analyses and log.analyses[0].incident_group_id
    }
    issue_by_group: dict = {}
    if group_ids:
        rows = (
            db.query(LogAnalysis.incident_group_id, GitlabIssue)
            .join(IncidentDecision, IncidentDecision.analysis_id == LogAnalysis.id)
            .join(GitlabIssue, GitlabIssue.incident_decision_id == IncidentDecision.id)
            .filter(LogAnalysis.incident_group_id.in_(group_ids))
            .all()
        )
        for group_id, issue in rows:
            issue_by_group.setdefault(group_id, issue)

    results = []
    for log in logs:
        response = LogResponse.model_validate(log)
        analysis = log.analyses[0] if log.analyses else None
        group_id = analysis.incident_group_id if analysis else None
        existing_issue = issue_by_group.get(group_id) if group_id else None
        response.gitlab_issue_url = existing_issue.web_url if existing_issue else None
        response.can_create_gitlab_issue = gitlab_configured and bool(log.decisions) and existing_issue is None
        results.append(response)
    return results


@router.get("", response_model=List[LogResponse])
def get_logs(
    service_name: Optional[str] = None,
    instance_id: Optional[str] = None,
    log_level: Optional[str] = None,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
    search_query: Optional[str] = None,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db)
):
    """
    Retrieve logs with optional filters and pagination.
    """
    logs = log_repo.get_all(
        db,
        skip=skip,
        limit=limit,
        service_name=service_name,
        instance_id=instance_id,
        log_level=log_level,
        start_date=start_date,
        end_date=end_date,
        search_query=search_query
    )
    return _attach_gitlab_issue_info(db, logs)

@router.get("/search", response_model=List[LogResponse])
def search_logs(
    q: str = Query(..., min_length=1),
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db)
):
    """
    Search log messages and stack traces matching a string.
    """
    logs = log_repo.get_all(
        db,
        skip=skip,
        limit=limit,
        search_query=q
    )
    return logs

@router.get("/summary", response_model=LogSummaryResponse)
def get_logs_summary(
    service_name: Optional[str] = None,
    instance_id: Optional[str] = None,
    log_level: Optional[str] = None,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
    search_query: Optional[str] = None,
    db: Session = Depends(get_db)
):
    """
    Aggregated view (level breakdown, top services, time series) of every log
    matching the given filters -- powers the Log Explorer's chart panel.
    """
    filters = dict(
        service_name=service_name,
        instance_id=instance_id,
        log_level=log_level,
        start_date=start_date,
        end_date=end_date,
        search_query=search_query
    )

    granularity = "hour" if start_date and end_date and (end_date - start_date) <= timedelta(days=2) else "day"

    total_matched = log_repo.count_matching(db, **filters)
    level_counts = log_repo.get_level_counts(db, **filters)
    top_services = log_repo.get_top_services(db, limit=8, **filters)
    time_series = log_repo.get_time_series(db, bucket_expr=_bucket_expr(granularity), **filters)
    for point in time_series:
        point["bucket"] = _parse_bucket(point["bucket"], granularity)
    errors_with_rca = log_repo.count_errors_with_rca(db, **filters)

    return LogSummaryResponse(
        total_matched=total_matched,
        level_counts=level_counts,
        top_services=top_services,
        time_series=time_series,
        bucket_granularity=granularity,
        errors_total=level_counts.get("ERROR", 0),
        errors_with_rca=errors_with_rca
    )

_EXPORT_COLUMNS = [
    "timestamp", "log_level", "service_name", "instance_id", "message", "stacktrace",
    "file_name", "file_path", "rca_solution_found", "rca_severity", "rca_root_cause_category"
]

@router.get("/export")
def export_logs(
    service_name: Optional[str] = None,
    instance_id: Optional[str] = None,
    log_level: Optional[str] = None,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
    search_query: Optional[str] = None,
    db: Session = Depends(get_db)
):
    """
    Exports the currently filtered logs as an Excel workbook, capped at
    EXPORT_ROW_LIMIT rows (same ordering as the table: newest first).
    ERROR rows are highlighted -- green if an RCA solution has already been
    generated, red if not -- so unresolved errors stand out at a glance.
    Total-matched vs rows-exported are surfaced via headers so the UI can
    flag a truncated export.
    """
    filters = dict(
        service_name=service_name,
        instance_id=instance_id,
        log_level=log_level,
        start_date=start_date,
        end_date=end_date,
        search_query=search_query
    )

    total_matched = log_repo.count_matching(db, **filters)
    rows = log_repo.get_for_export(db, limit=EXPORT_ROW_LIMIT, **filters)

    wb = Workbook()
    ws = wb.active
    ws.title = "Logs"
    ws.append(_EXPORT_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = _HEADER_FILL

    for log in rows:
        ws.append([
            log.timestamp.isoformat(),
            log.log_level,
            log.service_name,
            log.instance_id or "",
            _excel_safe(log.message),
            _excel_safe(log.stacktrace),
            log.file_name,
            log.file_path,
            "Yes" if log.has_rca else "No",
            log.rca_severity or "",
            log.rca_root_cause_category or "",
        ])

        if log.log_level == "ERROR":
            fill = _SOLVED_FILL if log.has_rca else _UNSOLVED_FILL
            for col in range(1, len(_EXPORT_COLUMNS) + 1):
                ws.cell(row=ws.max_row, column=col).fill = fill

    for col in range(1, len(_EXPORT_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22
    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"logs_export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    headers = {
        "Content-Disposition": f'attachment; filename="{filename}"',
        "X-Total-Matched": str(total_matched),
        "X-Rows-Exported": str(len(rows)),
    }
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers
    )

@router.get("/{log_id}", response_model=LogResponse)
def get_log_by_id(log_id: int, db: Session = Depends(get_db)):
    """
    Fetch a single log record by ID.
    """
    log = log_repo.get(db, log_id)
    if not log:
        raise HTTPException(status_code=404, detail="Log record not found")
    return log

def _reprocess_single_file(db: Session, db_file) -> tuple[bool, str]:
    """
    Fully re-reads and re-parses a tracked file from scratch. Existing rows for
    this file are only deleted *after* the file has been successfully read and
    parsed -- never before -- so a network share being briefly unreachable (or
    any other read/parse failure) leaves prior data intact instead of wiping it.
    """
    file_path = db_file.file_path

    if not os.path.isfile(file_path):
        return False, "file not found or unreachable (e.g. network share down)"
    if not os.access(file_path, os.R_OK):
        return False, "file not readable (permission denied)"

    try:
        service_name, instance_id = parse_source_identity(file_path, get_active_root_paths(db))
        with open(file_path, "rb") as f:
            content_bytes = f.read()
        content = content_bytes.decode("utf-8", errors="replace")
        lines = content.splitlines()
        parsed_logs = log_parser.parse_lines(
            lines, db_file.file_name, file_path,
            default_service=service_name, instance_id=instance_id
        )
    except Exception as e:
        logger.error(f"Reprocess read/parse failed for {file_path}: {str(e)}", exc_info=True)
        return False, f"read/parse error: {str(e)}"

    # Only now, having confirmed the file is fully readable and parseable, do
    # we touch existing data for it.
    log_repo.delete_by_path(db, file_path)
    if parsed_logs:
        created = log_repo.create_many(db, parsed_logs)
        for original, db_obj in zip(parsed_logs, created):
            auto_trigger.maybe_schedule_auto_rca(db_obj.id, original.log_level)

    db_file.last_processed_position = len(content_bytes)
    db_file.last_processed_time = datetime.utcnow()
    db_file.status = "completed"
    db.add(db_file)
    db.commit()
    return True, "ok"

@router.post("/reprocess")
def reprocess_logs(
    request: Optional[LogReprocessRequest] = None,
    db: Session = Depends(get_db)
):
    """
    Reprocesses logs for a specific file or all files. See
    _reprocess_single_file for the delete-after-verify safety guarantee.
    """
    request_data = request or LogReprocessRequest()
    files_to_reprocess = []

    if request_data.file_id:
        db_file = log_file_repo.get(db, request_data.file_id)
        if not db_file:
            raise HTTPException(status_code=404, detail="Log source file not found")
        files_to_reprocess.append(db_file)
    elif request_data.file_path:
        db_file = log_file_repo.get_by_path(db, request_data.file_path)
        if not db_file:
            raise HTTPException(status_code=404, detail="Log source file not found")
        files_to_reprocess.append(db_file)
    else:
        # Reprocess all files
        # Synchronously scan directory first to register any newly placed log files (crucial for SQLite environments)
        log_watcher_service.scan_directory()
        files_to_reprocess = log_file_repo.get_all(db)

    reprocessed = []
    skipped = []
    for db_file in files_to_reprocess:
        ok, reason = _reprocess_single_file(db, db_file)
        if ok:
            reprocessed.append(db_file.file_path)
        else:
            skipped.append({"file_path": db_file.file_path, "reason": reason})

    message = f"Reprocessed {len(reprocessed)} file(s)."
    if skipped:
        message += f" Skipped {len(skipped)} unreachable/unreadable file(s) -- their existing data was preserved."

    return {"message": message, "reprocessed": reprocessed, "skipped": skipped}
